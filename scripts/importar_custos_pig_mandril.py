"""Importação inicial do módulo GESTÃO DE CUSTOS — Fase 1 (grupo PIG MANDRIL).

NÃO é o que roda em produção: a importação real acontece sozinha, uma única
vez, no primeiro boot do app depois deste deploy (`_seed_custos_pig_mandril`
em app.py, guardado por ControleSistema — mesmo padrão de
`_seed_lead_time_transportadora`), lendo `data/custo_de_producao_20_09_2026.xlsx`
(a mesma planilha, já incluída no repositório). Este script aqui é a versão
de referência/iteração usada durante o desenvolvimento pra validar a lógica
localmente antes de portar pro app.py — mantido no repo como documentação
do processo e pra facilitar auditoria futura das fórmulas mapeadas (não em
texto explicativo da planilha, que tinha pelo menos 2 discrepâncias
confirmadas contra a fórmula real).

Rodar com: DATABASE_URL=... python3 importar_custos_pig_mandril.py
(precisa do caminho da planilha original — XLSX_PATH abaixo aponta pro
upload original desta sessão, ajuste se for rodar de novo depois)
"""
import os
import sys

sys.path.insert(0, "/home/claude/producao_app")
os.environ.setdefault("DATABASE_URL", "sqlite:////home/claude/producao_app/instance/pedidos.db")

import openpyxl

from app import create_app
from extensions import db
from models import EstruturaProduto, EstruturaProdutoItem, MateriaPrima, Produto

XLSX_PATH = "/root/.claude/uploads/f35f9898-3d96-5c11-807a-4ceee7b09718/61d8de41-CUSTO_DE_PRODUCAO__.xlsx"

HH_RATE_SEED = 40.0  # valor vigente na planilha na data da importação — usado só pra DERIVAR ciclo_horas


def _dn_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _dn_num_str(v):
    """Extrai só a parte numérica do DN, descartando aspas/polegadas (ex.: "6''" -> "6").

    Usado nas tabelas de PARÂMETROS de PIGS EM BORRACHA (linhas 65-69/73-77), cujo DN vem
    formatado como texto com aspas duplas de polegada (`6''`), diferente da tabela de
    tubo/flange (linhas 83-101) e das abas de produto (LBD/LUN), que usam DN numérico puro —
    sem esta normalização as duas tabelas nunca batem por chave.
    """
    if v is None:
        return None
    s = "".join(ch for ch in str(v).strip() if ch.isdigit() or ch == ".")
    if not s:
        return None
    try:
        f = float(s)
        return str(int(f)) if f == int(f) else str(f)
    except ValueError:
        return s


def _get_or_create_mp(cache, codigo, descricao, unidade, custo, categoria):
    mp = cache.get(codigo)
    if mp is not None:
        return mp
    mp = MateriaPrima.query.filter_by(codigo=codigo).first()
    if mp is None:
        mp = MateriaPrima(codigo=codigo, descricao=descricao, unidade=unidade, custo_atual=custo or 0, categoria=categoria, ativo=True)
        db.session.add(mp)
        db.session.flush()
    cache[codigo] = mp
    return mp


def _get_or_create_produto(cache, familia, codigo, descricao=None, categoria=None, chave_busca=None):
    key = (familia, codigo)
    p = cache.get(key)
    if p is not None:
        return p
    p = Produto.query.filter_by(familia=familia, codigo=codigo).first()
    if p is None:
        p = Produto(familia=familia, codigo=codigo, descricao=descricao, categoria=categoria, chave_busca=chave_busca, ativo=True)
        db.session.add(p)
        db.session.flush()
    else:
        # script idempotente: mantém o cadastro sincronizado com a última versão do import
        # (ex.: chave_busca adicionada numa correção posterior a um produto já existente)
        if descricao is not None:
            p.descricao = descricao
        if categoria is not None:
            p.categoria = categoria
        if chave_busca is not None:
            p.chave_busca = chave_busca
    cache[key] = p
    return p


def _get_or_create_estrutura(produto, dn, ciclo_horas):
    e = EstruturaProduto.query.filter_by(produto_id=produto.id, dn=dn).first()
    if e is None:
        e = EstruturaProduto(produto_id=produto.id, dn=dn, ciclo_horas=ciclo_horas or 0, ativo=True)
        db.session.add(e)
        db.session.flush()
    else:
        e.ciclo_horas = ciclo_horas or 0
        # zera os itens antigos pra reimportar limpo (script idempotente)
        for item in list(e.itens):
            db.session.delete(item)
        db.session.flush()
    return e


def _add_item(estrutura, ordem, tipo, quantidade, materia_prima=None, subproduto=None, observacao=None):
    if not quantidade:
        return
    db.session.add(EstruturaProdutoItem(
        estrutura_id=estrutura.id, tipo=tipo, quantidade=quantidade,
        materia_prima_id=materia_prima.id if materia_prima else None,
        subproduto_id=subproduto.id if subproduto else None,
        observacao=observacao, ordem=ordem,
    ))


def main():
    app = create_app()
    with app.app_context():
        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
        mp_cache = {}
        produto_cache = {}

        # ------------------------------------------------------------
        # 1. Matérias-primas "químicas" — PARÂMETROS linhas 6-15
        # ------------------------------------------------------------
        ws_param = wb["PARÂMETROS"]
        quimicas_por_linha = {}  # linha -> MateriaPrima
        for r in range(6, 16):
            desc = ws_param.cell(r, 2).value
            preco = ws_param.cell(r, 5).value
            if not desc:
                continue
            codigo = "QUIM-" + "".join(ch for ch in desc.upper() if ch.isalnum())[:30]
            mp = _get_or_create_mp(mp_cache, codigo, desc, "kg", preco, "Química")
            quimicas_por_linha[r] = mp
        # mapeamento fixo linha->uso (confirmado nas fórmulas de PU CAST: $E$6, $E$8, $E$10, $E$11, $E$12)
        MP_PRE_TDI = quimicas_por_linha[6]       # 12-70 A (PRE, sistema TDI)
        MP_MOCA = quimicas_por_linha[8]          # MOCA CURATIVO TDI
        MP_PRE_A_18Q = quimicas_por_linha[10]    # 18Q (PRE "A", sistema MDI)
        MP_PRE_B_GS234 = quimicas_por_linha[11]  # IMUCURE GS234-960 (PRE "B", sistema MDI)
        MP_BUTANODIOL = quimicas_por_linha[12]   # IMUCURE GR45 1,4 BUTANODIOL

        # ------------------------------------------------------------
        # 2. Matérias-primas por DN — PARÂMETROS linhas 82-101 (LBD_REV A:
        #    tubo/flange bumper/flange solda/parafuso/arruela/porca)
        # ------------------------------------------------------------
        tubo_por_dn, flange_bumper_por_dn, flange_solda_por_dn = {}, {}, {}
        parafuso_por_dn, arruela_por_dn, porca_por_dn = {}, {}, {}
        for r in range(83, 102):
            dn = _dn_str(ws_param.cell(r, 1).value)
            if dn is None:
                continue
            tubo_por_dn[dn] = _get_or_create_mp(mp_cache, f"TUBO-DN{dn}", f"Tubo DN {dn}", "un", ws_param.cell(r, 2).value, "Componente DN")
            flange_bumper_por_dn[dn] = _get_or_create_mp(mp_cache, f"FLANGE-BUMPER-DN{dn}", f"Flange bumper DN {dn}", "un", ws_param.cell(r, 3).value, "Componente DN")
            flange_solda_por_dn[dn] = _get_or_create_mp(mp_cache, f"FLANGE-SOLDA-DN{dn}", f"Flange solda DN {dn}", "un", ws_param.cell(r, 4).value, "Componente DN")
            parafuso_por_dn[dn] = _get_or_create_mp(mp_cache, f"PARAFUSO-DN{dn}", f"Parafuso DN {dn}", "un", ws_param.cell(r, 5).value, "Componente DN")
            arruela_por_dn[dn] = _get_or_create_mp(mp_cache, f"ARRUELA-DN{dn}", f"Arruela DN {dn}", "un", ws_param.cell(r, 6).value, "Componente DN")
            porca_por_dn[dn] = _get_or_create_mp(mp_cache, f"PORCA-DN{dn}", f"Porca DN {dn}", "un", ws_param.cell(r, 7).value, "Componente DN")

        # ------------------------------------------------------------
        # 3. Matérias-primas PIGS EM BORRACHA — PARÂMETROS linhas 64-69
        #    (copo por material) e 72-77 (componentes comuns)
        # ------------------------------------------------------------
        copo_epdm_por_dn, copo_buna_por_dn, copo_viton_por_dn = {}, {}, {}
        for r in range(65, 70):
            dn = _dn_num_str(ws_param.cell(r, 1).value)
            if dn is None:
                continue
            copo_epdm_por_dn[dn] = _get_or_create_mp(mp_cache, f"COPO-BORRACHA-EPDM-DN{dn}", f"Copo de borracha EPDM DN {dn}", "un", ws_param.cell(r, 2).value, "Componente DN")
            copo_buna_por_dn[dn] = _get_or_create_mp(mp_cache, f"COPO-BORRACHA-BUNA-DN{dn}", f"Copo de borracha BUNA N DN {dn}", "un", ws_param.cell(r, 3).value, "Componente DN")
            copo_viton_por_dn[dn] = _get_or_create_mp(mp_cache, f"COPO-BORRACHA-VITON-DN{dn}", f"Copo de borracha VITON DN {dn}", "un", ws_param.cell(r, 4).value, "Componente DN")

        eixo_por_dn, cabecote_por_dn, nylon_por_dn, porca_bor_por_dn, flange_bor_por_dn = {}, {}, {}, {}, {}
        hh_montagem_borracha = None
        for r in range(73, 78):
            dn = _dn_num_str(ws_param.cell(r, 1).value)
            if dn is None:
                continue
            eixo_por_dn[dn] = _get_or_create_mp(mp_cache, f"EIXO-BORRACHA-DN{dn}", f"Eixo (barra roscada) DN {dn}", "un", ws_param.cell(r, 2).value, "Componente DN")
            cabecote_por_dn[dn] = _get_or_create_mp(mp_cache, f"CABECOTE-BORRACHA-DN{dn}", f"Cabeçote PU DN {dn}", "un", ws_param.cell(r, 3).value, "Componente DN")
            nylon_por_dn[dn] = _get_or_create_mp(mp_cache, f"NYLON-BORRACHA-DN{dn}", f"De nylon preto (2X) DN {dn}", "un", ws_param.cell(r, 4).value, "Componente DN")
            porca_bor_por_dn[dn] = _get_or_create_mp(mp_cache, f"PORCA-BORRACHA-DN{dn}", f"Porca DN {dn}", "un", ws_param.cell(r, 5).value, "Componente DN")
            flange_bor_por_dn[dn] = _get_or_create_mp(mp_cache, f"FLANGE-BORRACHA-DN{dn}", f"Flange (2X) DN {dn}", "un", ws_param.cell(r, 6).value, "Componente DN")
            hh_montagem_borracha = ws_param.cell(r, 7).value  # 120 pra todas as DNs

        # ------------------------------------------------------------
        # 4. Família PU CAST — DS, DG, DE, COPO CONICO, COPO PISTAO,
        #    HFLEX, DISCFLEX SD, DISCFLEX SDI.
        #
        #    IMPORTANTE (achado na verificação cruzada contra BUSCA DE
        #    CUSTO): a coluna "CUSTO MP" (T/AC) de cada linha NÃO é sempre
        #    peso×preço dos 5 componentes químicos (E/G/I/K/M) somados —
        #    em vários casos é só F+H (2 dos 5), às vezes com F referenciando
        #    um insumo químico DIFERENTE do que outras linhas do mesmo tipo
        #    de item usam (ex.: DS linha 10 usa PARÂMETROS!$E$6, DE linha 59
        #    usa PARÂMETROS!$E$14 — não é uma correspondência fixa por
        #    coluna). Decompor por química ficaria errado sem reler a
        #    fórmula de CADA linha individualmente. Por isso, aqui a
        #    matéria-prima entra como 1 valor já consolidado por (item, DN)
        #    — mesmo critério já usado em ELC_MG_PC — preservando fidelidade
        #    exata ao valor da planilha (validado 1:1 contra BUSCA DE CUSTO)
        #    em vez de arriscar uma decomposição incorreta. "BUMPER (2X)"
        #    ficou de fora (a própria aba referencia a LBD para esse valor,
        #    não tem BOM própria).
        # ------------------------------------------------------------
        ws_pu = wb["PU CAST"]
        col_map = {"ITEM": 27, "CUSTO_MP": 29, "CUSTO_HH_UNIT": 34, "DN_NUM": 37}
        # chave de busca pra casar com o texto livre dos pedidos do PCP — confirmado por
        # amostragem real em instance/pedidos.db (ex. "DISCO SELO DN 12"", "PIG DISCFLEX SDI DN 6"")
        # esses itens são vendidos/produzidos como sobressalente avulso, não só dentro da LBD/LUN.
        CHAVE_BUSCA_PU_CAST = {
            "DS": "DISCO SELO",
            "DG": "DISCO GUIA",
            "DE": "DISCO ESPAÇADOR",
            "COPO CONICO": "COPO CONICO",
            "COPO PISTAO": "COPO PIST",
            "HFLEX": "HFLEX",
            "DISCFLEX SD": "DISCFLEX SD",
            "DISCFLEX SDI": "DISCFLEX SDI",
        }
        pu_cast_estruturas = {}  # (item, dn_str) -> EstruturaProduto, pra LBD/LUN referenciarem
        n_pu = 0
        for r in range(10, ws_pu.max_row + 1):
            item = ws_pu.cell(r, col_map["ITEM"]).value
            dn_num = ws_pu.cell(r, col_map["DN_NUM"]).value
            if item is None or item == "BUMPER (2X)" or dn_num in (None, ""):
                continue
            dn = _dn_str(dn_num)
            custo_mp = ws_pu.cell(r, col_map["CUSTO_MP"]).value or 0
            custo_hh_unit = ws_pu.cell(r, col_map["CUSTO_HH_UNIT"]).value or 0
            ciclo_horas = round((custo_hh_unit / HH_RATE_SEED), 6) if custo_hh_unit else 0

            produto = _get_or_create_produto(
                produto_cache, "PU CAST", item, descricao=f"PU CAST — {item}", categoria="Sobressalente",
                chave_busca=CHAVE_BUSCA_PU_CAST.get(item),
            )
            estrutura = _get_or_create_estrutura(produto, dn, ciclo_horas)

            slug = "".join(ch for ch in str(item).upper() if ch.isalnum())[:20]
            mp = _get_or_create_mp(
                mp_cache, f"PUCAST-MP-{slug}-DN{dn}", f"PU CAST {item} — matéria-prima consolidada DN {dn}",
                "un", custo_mp, "Química (consolidada)",
            )
            _add_item(estrutura, 0, "MATERIA_PRIMA", 1, materia_prima=mp)
            pu_cast_estruturas[(item, dn)] = estrutura
            n_pu += 1
        print(f"PU CAST: {n_pu} estruturas (produto, DN) importadas.")

        # ------------------------------------------------------------
        # 5. Família LBD — produto único "LBD-DG2-DS4"
        # ------------------------------------------------------------
        ws_lbd = wb["LBD"]
        produto_lbd = _get_or_create_produto(produto_cache, "LBD", "LBD-DG2-DS4", descricao="Mandril LBD-DG2-DS4", categoria="PIG", chave_busca="LBD")
        n_lbd = 0
        for r in range(6, 28):
            modelo = ws_lbd.cell(r, 1).value
            dn_val = ws_lbd.cell(r, 2).value
            if modelo is None or dn_val is None:
                continue
            dn = _dn_str(dn_val)
            custo_hh_unit = ws_lbd.cell(r, 59).value or 0  # BG
            ciclo_horas = round(custo_hh_unit / HH_RATE_SEED, 6) if custo_hh_unit else 0
            estrutura = _get_or_create_estrutura(produto_lbd, dn, ciclo_horas)

            ordem = 0
            # discos (subproduto de PU CAST) — quantidade = multiplicador fixo já confirmado nas fórmulas (F9=E9*4, J9=I9*2, N9=M9*6)
            for item_pu, qtd in (("DS", 4), ("DG", 2), ("DE", 6)):
                sub = pu_cast_estruturas.get((item_pu, dn))
                if sub is not None:
                    _add_item(estrutura, ordem, "SUBPRODUTO", qtd, subproduto=produto_cache[("PU CAST", item_pu)])
                    ordem += 1
            # tubo/flanges/fixação (matéria-prima por DN)
            tubo_qtd = ws_lbd.cell(r, 17).value or 0            # Q
            flange_bumper_qtd = 2 if ws_lbd.cell(r, 24).value else 0   # Y = X*2 (2 unidades)
            bumper_peso_total = ws_lbd.cell(r, 27).value or 0   # AA (peso total elastômero, kg)
            flange_solda_qtd = 2 if ws_lbd.cell(r, 35).value else 0    # AI = AH*2
            parafuso_qtd = ws_lbd.cell(r, 38).value or 0        # AL
            arruela_qtd = ws_lbd.cell(r, 43).value or 0         # AQ
            porca_qtd = ws_lbd.cell(r, 48).value or 0           # AV

            if tubo_qtd and dn in tubo_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", tubo_qtd, materia_prima=tubo_por_dn[dn]); ordem += 1
            if flange_bumper_qtd and dn in flange_bumper_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_bumper_qtd, materia_prima=flange_bumper_por_dn[dn]); ordem += 1
            if bumper_peso_total:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", bumper_peso_total, materia_prima=MP_PRE_TDI, observacao="bumper (elastômero)"); ordem += 1
            if flange_solda_qtd and dn in flange_solda_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_solda_qtd, materia_prima=flange_solda_por_dn[dn]); ordem += 1
            if parafuso_qtd and dn in parafuso_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", parafuso_qtd, materia_prima=parafuso_por_dn[dn]); ordem += 1
            if arruela_qtd and dn in arruela_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", arruela_qtd, materia_prima=arruela_por_dn[dn]); ordem += 1
            if porca_qtd and dn in porca_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", porca_qtd, materia_prima=porca_por_dn[dn]); ordem += 1
            n_lbd += 1
        print(f"LBD: {n_lbd} DNs importadas.")

        # ------------------------------------------------------------
        # 6. Família LUN — produto único "LUN" (mesma estrutura da LBD,
        #    mas com quantidades de disco variáveis por DN, lidas direto
        #    da planilha em vez de multiplicador fixo)
        # ------------------------------------------------------------
        ws_lun = wb["LUN"]
        produto_lun = _get_or_create_produto(produto_cache, "LUN", "LUN", descricao="Mandril LUN", categoria="PIG", chave_busca="LUN")
        n_lun = 0
        for r in range(6, 28):
            modelo = ws_lun.cell(r, 1).value
            dn_val = ws_lun.cell(r, 2).value
            if modelo is None or dn_val is None:
                continue
            dn = _dn_str(dn_val)
            custo_hh_unit = ws_lun.cell(r, 59).value or 0  # BG
            ciclo_horas = round(custo_hh_unit / HH_RATE_SEED, 6) if custo_hh_unit else 0
            estrutura = _get_or_create_estrutura(produto_lun, dn, ciclo_horas)

            ordem = 0
            for item_pu, qtd_col in (("COPO CONICO", 3), ("DS", 6), ("DG", 9), ("DE", 12)):  # C, F, I, L
                qtd = ws_lun.cell(r, qtd_col).value or 0
                sub = pu_cast_estruturas.get((item_pu, dn))
                if qtd and sub is not None:
                    _add_item(estrutura, ordem, "SUBPRODUTO", qtd, subproduto=produto_cache[("PU CAST", item_pu)])
                    ordem += 1

            tubo_qtd = ws_lun.cell(r, 17).value or 0
            flange_bumper_qtd = 2 if ws_lun.cell(r, 24).value else 0
            bumper_peso_total = ws_lun.cell(r, 27).value or 0
            flange_solda_qtd = 2 if ws_lun.cell(r, 35).value else 0
            parafuso_qtd = ws_lun.cell(r, 38).value or 0
            arruela_qtd = ws_lun.cell(r, 43).value or 0
            porca_qtd = ws_lun.cell(r, 48).value or 0

            if tubo_qtd and dn in tubo_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", tubo_qtd, materia_prima=tubo_por_dn[dn]); ordem += 1
            if flange_bumper_qtd and dn in flange_bumper_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_bumper_qtd, materia_prima=flange_bumper_por_dn[dn]); ordem += 1
            if bumper_peso_total:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", bumper_peso_total, materia_prima=MP_PRE_TDI, observacao="bumper (elastômero)"); ordem += 1
            if flange_solda_qtd and dn in flange_solda_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_solda_qtd, materia_prima=flange_solda_por_dn[dn]); ordem += 1
            if parafuso_qtd and dn in parafuso_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", parafuso_qtd, materia_prima=parafuso_por_dn[dn]); ordem += 1
            if arruela_qtd and dn in arruela_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", arruela_qtd, materia_prima=arruela_por_dn[dn]); ordem += 1
            if porca_qtd and dn in porca_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", porca_qtd, materia_prima=porca_por_dn[dn]); ordem += 1
            n_lun += 1
        print(f"LUN: {n_lun} DNs importadas.")

        # ------------------------------------------------------------
        # 7. Família CORPO MANDRIL — "corpo + fixação" (sem os discos).
        #
        #    Confirmado via fórmula real: CORPO MANDRIL!D = LBD!AZ =
        #    LBD!S+Y+AC+AI (tubo + flange bumper + bumper elastômero +
        #    flange solda) e CORPO MANDRIL!E = LBD!BA = LBD!AN+AS+AX
        #    (parafuso + arruela + porca) — SEMPRE lidas da MESMA LINHA
        #    da LBD (mesma DN), usando as quantidades REAIS daquela linha
        #    (que variam por DN — não são 12/2/2 fixos como uma primeira
        #    leitura sugeriria). Por isso este bloco relê a aba LBD (não a
        #    aba CORPO MANDRIL, que só tem o DN e o total, sem as
        #    quantidades) com a mesma extração de colunas já usada na
        #    seção 5, mas SEM os discos DS/DG/DE. A aba CORPO MANDRIL
        #    ainda é usada só para o HH/ciclo (coluna H, independente).
        #    Verificado 1:1 contra CORPO MANDRIL!I (coluna "CUSTO TOTAL")
        #    após a correção — bate exatamente (as 3 primeiras linhas,
        #    DN 2/3/4, ficam com custo 0 porque a LBD não tem fixação
        #    cadastrada pra esses DNs na planilha original).
        # ------------------------------------------------------------
        ws_corpo = wb["CORPO MANDRIL"]
        produto_corpo = _get_or_create_produto(produto_cache, "CORPO MANDRIL", "CORPO + FIXAÇÃO", descricao="Corpo do PIG (tubo+flanges+fixação, sem discos)", categoria="Sobressalente", chave_busca="CORPO MANDRIL")
        n_corpo = 0
        for r in range(6, 28):
            dn_val = ws_lbd.cell(r, 2).value
            if dn_val is None:
                continue
            dn = _dn_str(dn_val)

            # HH/ciclo — lido da própria aba CORPO MANDRIL (coluna H = CUSTO HH CICLO, lote inteiro, qnt=1 pç)
            custo_hh_cico = 0
            for rc in range(4, ws_corpo.max_row + 1):
                if _dn_str(ws_corpo.cell(rc, 3).value) == dn:
                    custo_hh_cico = ws_corpo.cell(rc, 8).value or 0
                    break
            ciclo_horas = round(custo_hh_cico / HH_RATE_SEED, 6) if custo_hh_cico else 0
            estrutura = _get_or_create_estrutura(produto_corpo, dn, ciclo_horas)

            ordem = 0
            tubo_qtd = ws_lbd.cell(r, 17).value or 0
            flange_bumper_qtd = 2 if ws_lbd.cell(r, 24).value else 0
            bumper_peso_total = ws_lbd.cell(r, 27).value or 0
            flange_solda_qtd = 2 if ws_lbd.cell(r, 35).value else 0
            parafuso_qtd = ws_lbd.cell(r, 38).value or 0
            arruela_qtd = ws_lbd.cell(r, 43).value or 0
            porca_qtd = ws_lbd.cell(r, 48).value or 0

            if tubo_qtd and dn in tubo_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", tubo_qtd, materia_prima=tubo_por_dn[dn]); ordem += 1
            if flange_bumper_qtd and dn in flange_bumper_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_bumper_qtd, materia_prima=flange_bumper_por_dn[dn]); ordem += 1
            if bumper_peso_total:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", bumper_peso_total, materia_prima=MP_PRE_TDI, observacao="bumper (elastômero)"); ordem += 1
            if flange_solda_qtd and dn in flange_solda_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", flange_solda_qtd, materia_prima=flange_solda_por_dn[dn]); ordem += 1
            if parafuso_qtd and dn in parafuso_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", parafuso_qtd, materia_prima=parafuso_por_dn[dn]); ordem += 1
            if arruela_qtd and dn in arruela_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", arruela_qtd, materia_prima=arruela_por_dn[dn]); ordem += 1
            if porca_qtd and dn in porca_por_dn:
                _add_item(estrutura, ordem, "MATERIA_PRIMA", porca_qtd, materia_prima=porca_por_dn[dn]); ordem += 1
            n_corpo += 1
        print(f"CORPO MANDRIL: {n_corpo} DNs importadas.")

        # ------------------------------------------------------------
        # 8. Família ELC_MG_PC — ELC (AÇO), ELP (PP), CINTA MAGNÉTICA,
        #    PLACA CALIBRADORA — custo MP já vem pronto por DN (fonte
        #    externa não enviada), vira 1 matéria-prima por (item, DN).
        # ------------------------------------------------------------
        ws_elc = wb["ELC_MG_PC"]
        n_elc = 0
        for r in range(4, ws_elc.max_row + 1):
            item = ws_elc.cell(r, 2).value
            dn_val = ws_elc.cell(r, 3).value
            custo_mp = ws_elc.cell(r, 4).value
            custo_hh_unit = ws_elc.cell(r, 9).value or 0
            if item is None or dn_val is None:
                continue
            dn = _dn_str(dn_val)
            slug = "".join(ch for ch in item.upper() if ch.isalnum())[:20]
            mp_codigo = f"{slug}-MP-DN{dn}"
            mp = _get_or_create_mp(mp_cache, mp_codigo, f"{item} — matéria-prima DN {dn}", "un", custo_mp, "Acessório")

            produto = _get_or_create_produto(produto_cache, "ELC_MG_PC", item, descricao=item, categoria="Acessório", chave_busca=item.split(" ")[0])
            ciclo_horas = round(custo_hh_unit / HH_RATE_SEED, 6) if custo_hh_unit else 0
            estrutura = _get_or_create_estrutura(produto, dn, ciclo_horas)
            _add_item(estrutura, 0, "MATERIA_PRIMA", 1, materia_prima=mp)
            n_elc += 1
        print(f"ELC_MG_PC: {n_elc} (item, DN) importados.")

        # ------------------------------------------------------------
        # 9. Família PIGS EM BORRACHA — 3 variantes de material (EPDM,
        #    BUNA N, VITON), DN 6/8/10/12/14
        # ------------------------------------------------------------
        ciclo_borracha = round((hh_montagem_borracha or 0) / HH_RATE_SEED, 6)
        variantes_borracha = (
            ("PIG LUN-CP3 EPDM", copo_epdm_por_dn, "LUN-CP3 EPDM"),
            ("PIG LUN-CP3 BUNA N", copo_buna_por_dn, "LUN-CP3 BUNA"),
            ("PIG LUN-CP3 VITON", copo_viton_por_dn, "LUN-CP3 VITON"),
        )
        n_borracha = 0
        for codigo, copo_por_dn, chave in variantes_borracha:
            produto = _get_or_create_produto(produto_cache, "PIGS EM BORRACHA", codigo, descricao=codigo, categoria="PIG", chave_busca=chave)
            for dn in ("6", "8", "10", "12", "14"):
                estrutura = _get_or_create_estrutura(produto, dn, ciclo_borracha)
                ordem = 0
                if dn in copo_por_dn:
                    # planilha PIGS EM BORRACHA linha "COPO BORRACHA ___ (3X)" — 3 copos por unidade
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 3, materia_prima=copo_por_dn[dn], observacao="3X"); ordem += 1
                if dn in eixo_por_dn:
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 1, materia_prima=eixo_por_dn[dn]); ordem += 1
                if dn in cabecote_por_dn:
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 1, materia_prima=cabecote_por_dn[dn]); ordem += 1
                if dn in nylon_por_dn:
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 1, materia_prima=nylon_por_dn[dn]); ordem += 1
                if dn in porca_bor_por_dn:
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 1, materia_prima=porca_bor_por_dn[dn]); ordem += 1
                if dn in flange_bor_por_dn:
                    _add_item(estrutura, ordem, "MATERIA_PRIMA", 1, materia_prima=flange_bor_por_dn[dn]); ordem += 1
                n_borracha += 1
        print(f"PIGS EM BORRACHA: {n_borracha} (variante, DN) importados.")

        db.session.commit()
        print("Importação concluída e commitada.")
        print(f"Total MateriaPrima: {MateriaPrima.query.count()}")
        print(f"Total Produto: {Produto.query.count()}")
        print(f"Total EstruturaProduto: {EstruturaProduto.query.count()}")
        print(f"Total EstruturaProdutoItem: {EstruturaProdutoItem.query.count()}")


if __name__ == "__main__":
    main()
